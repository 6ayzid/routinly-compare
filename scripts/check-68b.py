import os
import openpyxl

REPORTS_PATH = r"C:\Users\bayzid\Documents\routinly-v1.5\reports\Routinly_Fall2026_Master_Routine.xlsx"
DOWNLOADS_PATH = r"C:\Users\bayzid\Downloads\Telegram Desktop\Routinly_Fall2026_Master_Routine.xlsx"
EXCEL_PATH = REPORTS_PATH if os.path.exists(REPORTS_PATH) else DOWNLOADS_PATH

wb = openpyxl.load_workbook(EXCEL_PATH, data_only=True)
sheet = wb.active

current_day = None
time_slots = {
    0: ('08:30', '10:00'),
    1: ('10:00', '11:30'),
    2: ('11:30', '13:00'),
    3: ('13:00', '14:30'),
    4: ('14:30', '16:00'),
    5: ('16:00', '17:30'),
}

days_order = ['SATURDAY', 'SUNDAY', 'MONDAY', 'TUESDAY', 'WEDNESDAY', 'THURSDAY']
entries = []

for r in range(1, sheet.max_row + 1):
    val_a = sheet.cell(r, 1).value
    if val_a and str(val_a).strip().upper() in days_order:
        current_day = str(val_a).strip().upper()
        continue
    
    # Check slots across columns: each slot has 3 columns: Room, Course, Teacher
    for slot_idx in range(6):
        c_room = slot_idx * 3 + 1
        c_course = slot_idx * 3 + 2
        c_teacher = slot_idx * 3 + 3
        
        room = sheet.cell(r, c_room).value
        course = sheet.cell(r, c_course).value
        teacher = sheet.cell(r, c_teacher).value
        
        if course and '68_B' in str(course):
            entries.append({
                'day': current_day,
                'slot': slot_idx + 1,
                'start_time': time_slots[slot_idx][0],
                'end_time': time_slots[slot_idx][1],
                'room': str(room).strip() if room else 'N/A',
                'course': str(course).strip(),
                'teacher': str(teacher).strip() if teacher else 'N/A',
                'excel_row': r,
            })

print(f'Total occurrences of 68_B in Excel: {len(entries)}')
for e in sorted(entries, key=lambda x: (days_order.index(x['day']), x['slot'])):
    time_str = f"{e['start_time']} - {e['end_time']}"
    print(f"{e['day']:<10} | Slot {e['slot']} ({time_str:<13}) | Room: {e['room']:<12} | Course: {e['course']:<15} | Teacher: {e['teacher']}")
